"""
Routes Documents
CRUD + Upload + Pipeline
"""
import math
import os
import tempfile
from pathlib import Path
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query, BackgroundTasks
from fastapi.responses import FileResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, and_, desc
from loguru import logger
from backend.database import get_db
from backend.models.document import Document, DocumentStatus, DocumentType
from backend.schemas.document import (
    DocumentResponse, DocumentUpdate, PipelineResult,
    UploadResponse, PaginatedDocuments
)
from backend.core.pipeline import get_pipeline
from backend.config import settings

router = APIRouter(prefix="/documents", tags=["Documents"])


# Upload
@router.post("/upload", response_model=UploadResponse, summary="Upload + Pipeline IA")
async def upload_document(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    max_size = settings.max_upload_size_mb * 1024 * 1024
    content = await file.read()
    if len(content) > max_size:
        raise HTTPException(400, f"Fichier trop volumineux. Max: {settings.max_upload_size_mb}MB")
    if not file.filename:
        raise HTTPException(400, "Nom de fichier manquant")

    upload_dir = Path(settings.archive_root_path).parent / "uploads_temp"
    upload_dir.mkdir(parents=True, exist_ok=True)
    temp_path = upload_dir / file.filename

    try:
        with open(temp_path, "wb") as f:
            f.write(content)
        original_path = str(temp_path)
        logger.info(f"[Upload] Fichier temporaire sauvegarde : {temp_path}")
    except Exception as e:
        logger.error(f"[Upload] Impossible de sauvegarder le fichier temporaire : {e}")
        original_path = f"upload/{file.filename}"
        temp_path = None

    pipeline = get_pipeline()
    result = await pipeline.process_document(
        db=db,
        file_content=content,
        filename=file.filename,
        original_path=original_path,
        file_size_kb=len(content) / 1024,
    )

    if temp_path and temp_path.exists():
        try:
            temp_path.unlink()
        except Exception:
            pass

    if result.status.value == "error":
        raise HTTPException(422, f"Erreur pipeline: {result.errors}")

    from backend.utils.filename_generator import generate_filename_from_pipeline
    suggested = None
    if not result.is_duplicate:
        suggested = generate_filename_from_pipeline(result.model_dump(), file.filename)

    return UploadResponse(
        message=f"Document {'duplique' if result.is_duplicate else 'archive'} avec succes",
        document_id=result.document_id,
        pipeline_result=result,
        suggested_filename=suggested,
    )


@router.post("/upload/batch", summary="Upload en lot")
async def upload_batch(
    files: list[UploadFile] = File(...),
    background_tasks: BackgroundTasks = BackgroundTasks(),
    db: AsyncSession = Depends(get_db),
):
    if len(files) > 50:
        raise HTTPException(400, "Maximum 50 fichiers par lot")

    upload_dir = Path(settings.archive_root_path).parent / "uploads_temp"
    upload_dir.mkdir(parents=True, exist_ok=True)

    results = []
    pipeline = get_pipeline()
    temp_paths = []

    for file in files:
        content = await file.read()
        filename = file.filename or "unknown.pdf"

        temp_path = upload_dir / filename
        try:
            with open(temp_path, "wb") as f:
                f.write(content)
            original_path = str(temp_path)
            temp_paths.append(temp_path)
        except Exception:
            original_path = ""
            temp_path = None

        result = await pipeline.process_document(
            db=db,
            file_content=content,
            filename=filename,
            original_path=original_path,
            file_size_kb=len(content) / 1024,
        )
        results.append({
            "filename": filename,
            "status": result.status.value,
            "document_id": result.document_id,
            "is_duplicate": result.is_duplicate,
            "ner": result.ner.model_dump() if result.ner else {},
            "classification": result.classification.model_dump() if result.classification else {},
            "errors": result.errors,
        })

    for tp in temp_paths:
        try:
            if tp and tp.exists():
                tp.unlink()
        except Exception:
            pass

    success = sum(1 for r in results if r["status"] == "archived")
    return {
        "message": f"Lot traite: {success}/{len(files)} archives",
        "total": len(files),
        "success": success,
        "results": results,
    }

@router.get("/classification-stats")
async def classification_stats(db: AsyncSession = Depends(get_db)):
    total = await db.scalar(select(func.count()).select_from(Document))
    high_conf = await db.scalar(select(func.count()).select_from(Document).where(Document.class_confidence >= 0.8))
    low_conf = await db.scalar(select(func.count()).select_from(Document).where(Document.class_confidence < 0.5))
    needs_review = await db.scalar(select(func.count()).select_from(Document).where(Document.needs_review == True))
    duplicates = await db.scalar(select(func.count()).select_from(Document).where(Document.is_duplicate == True))
    avg_ocr = await db.scalar(select(func.avg(Document.ocr_confidence)).where(Document.ocr_confidence > 0))

    ocr_rows = await db.execute(
        select(Document.aircraft_registration,
               func.avg(Document.ocr_confidence).label("avg_ocr"),
               func.count().label("doc_count"))
        .where(Document.ocr_confidence > 0)
        .group_by(Document.aircraft_registration)
        .order_by(Document.aircraft_registration)
    )

    return {
        "total": total,
        "high_confidence": high_conf,
        "low_confidence": low_conf,
        "needs_review": needs_review,
        "duplicates": duplicates,
        "avg_ocr": round(avg_ocr or 0, 1),
        "ocr_by_aircraft": [
            {"registration": r.aircraft_registration, "avg_ocr": round(r.avg_ocr, 1), "doc_count": r.doc_count}
            for r in ocr_rows if r.aircraft_registration
        ],
    }
# CRUD
@router.get("/", response_model=PaginatedDocuments, summary="Lister les documents")
async def list_documents(
    aircraft: Optional[str] = Query(None),
    doc_type: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    ata_chapter: Optional[str] = Query(None),
    es_reference: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
is_critical: Optional[bool] = Query(None),
    needs_review: Optional[bool] = Query(None),
    no_aircraft: Optional[bool] = Query(None),
    page: int = Query(1, ge=1),
    size: int = Query(20, ge=1, le=500),
    sort: str = Query("created_at"),
    db: AsyncSession = Depends(get_db),
):
    conditions = []
    if aircraft:
        conditions.append(Document.aircraft_registration.ilike(f"%{aircraft}%"))
    if doc_type:
        conditions.append(Document.doc_type == doc_type)
    if category:
        conditions.append(Document.category.ilike(f"%{category}%"))
    if ata_chapter:
        conditions.append(Document.ata_chapter.ilike(f"%{ata_chapter}%"))
    if es_reference:
        conditions.append(Document.es_reference.ilike(f"%{es_reference}%"))
    if status:
        conditions.append(Document.status == status)
    if is_critical is not None:
        conditions.append(Document.is_critical == is_critical)
    if no_aircraft:
        conditions.append(Document.aircraft_registration.is_(None))

    count_q = select(func.count(Document.id))
    if conditions:
        count_q = count_q.where(and_(*conditions))
    total = await db.scalar(count_q) or 0

    q = select(Document)
    if conditions:
        q = q.where(and_(*conditions))
    q = q.order_by(desc(Document.created_at)).offset((page - 1) * size).limit(size)
    result = await db.execute(q)
    docs = result.scalars().all()

    return PaginatedDocuments(
        items=[DocumentResponse.model_validate(d) for d in docs],
        total=total,
        page=page,
        size=size,
        pages=math.ceil(total / size),
    )


# Export XLSX
@router.get("/export/xlsx", summary="Exporter tous les documents en Excel")
async def export_xlsx(db: AsyncSession = Depends(get_db)):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse
    from datetime import date

    result = await db.execute(select(Document).order_by(desc(Document.created_at)))
    docs = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Documents NouvelAir"

    bleu_marine = "1B3A5C"
    headers = ["ID", "Fichier", "Avion", "Type", "Categorie", "Ref.ES", "ATA", "OCR%", "Corrige", "Date"]
    widths  = [6, 50, 10, 15, 18, 12, 8, 8, 9, 12]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill      = PatternFill("solid", fgColor=bleu_marine)
        cell.font      = Font(color="FFFFFF", bold=True, size=10, name="Arial")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 20

    TYPE_COLORS = {
        "WORK_ORDER": "EBF5FB", "JOBCARD": "F4ECF7",
        "AD":         "FADBD8", "RCT":    "D5F5E3",
        "CERTIFICATE":"FEF9E7", "SB":     "FDEBD0",
        "DB_CHART":   "EAF2FF", "SPECS":  "F0F0F0",
    }

    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    rouge  = "FADBD8"
    blanc  = "FFFFFF"

    for row_idx, doc in enumerate(docs, 2):
        dt = (doc.doc_type or "").upper().replace(" ", "_")
        fill_color = TYPE_COLORS.get(dt, blanc)
        if doc.needs_review:
            fill_color = rouge

        row_data = [
            doc.id,
            doc.filename or "",
            doc.aircraft_registration or "",
            doc.doc_type or "",
            doc.category or "",
            doc.es_reference or "",
            doc.ata_chapter or "",
            round(doc.ocr_confidence, 1) if doc.ocr_confidence else 0,
            "Oui" if doc.manually_corrected else "Non",
            doc.created_at.strftime("%Y-%m-%d") if doc.created_at else "",
        ]

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.fill      = PatternFill("solid", fgColor=fill_color)
            cell.font      = Font(size=9, name="Arial")
            cell.alignment = Alignment(vertical="center")
            cell.border    = border

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="documents_NouvelAir_{date.today()}.xlsx"'}
    )


# Export XLSX — Documents à réviser uniquement
@router.get("/export/needs-review", summary="Exporter les documents à réviser")
async def export_needs_review(db: AsyncSession = Depends(get_db)):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse
    from datetime import date

    result = await db.execute(
        select(Document)
        .where(Document.needs_review == True)
        .order_by(desc(Document.ocr_confidence))
    )
    docs = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "A Réviser"

    bleu_marine = "1B3A5C"
    orange      = "E67E22"
    headers = ["ID", "Fichier", "Avion", "Type", "Categorie", "Ref.ES", "ATA", "OCR%", "Classifier%", "Corrige", "Date"]
    widths  = [6, 50, 10, 15, 18, 12, 8, 8, 12, 9, 12]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill      = PatternFill("solid", fgColor=bleu_marine)
        cell.font      = Font(color="FFFFFF", bold=True, size=10, name="Arial")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 20
    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_idx, doc in enumerate(docs, 2):
        row_data = [
            doc.id,
            doc.filename or "",
            doc.aircraft_registration or "",
            doc.doc_type or "",
            doc.category or "",
            doc.es_reference or "",
            doc.ata_chapter or "",
            round(doc.ocr_confidence, 1) if doc.ocr_confidence else 0,
            round(doc.classifier_confidence * 100, 1) if doc.classifier_confidence else 0,
            "Oui" if doc.manually_corrected else "Non",
            doc.created_at.strftime("%Y-%m-%d") if doc.created_at else "",
        ]
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.fill      = PatternFill("solid", fgColor="FEF9E7")
            cell.font      = Font(size=9, name="Arial")
            cell.alignment = Alignment(vertical="center")
            cell.border    = border

    # Ligne résumé en bas
    ws.append([])
    ws.append(["", f"Total : {len(docs)} documents à réviser", "", "", "", "", "", "", "", "", ""])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="a_reviser_NouvelAir_{date.today()}.xlsx"'}
    )

# Import XLSX (sync Excel -> BDD)
@router.post("/import/xlsx", summary="Importer corrections depuis Excel")
async def import_xlsx(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    import io
    from openpyxl import load_workbook

    content = await file.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active

    updated = 0
    errors  = []
    skipped = 0

    TYPE_MAP = {
        "work order": "WORK_ORDER", "workorder": "WORK_ORDER",
        "jobcard": "JOBCARD", "job card": "JOBCARD",
        "ad": "AD", "airworthiness directive": "AD",
        "sb": "SB", "service bulletin": "SB",
        "certificate": "CERTIFICATE", "rct": "RCT",
        "atl": "ATL", "amm": "AMM", "cmm": "CMM",
        "ipc": "IPC", "specs": "SPECS", "db_chart": "DB_CHART",
        "d&b chart": "DB_CHART", "defect report": "DEFECT_REPORT",
        "other": "OTHER",
    }

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        try:
            doc_id   = int(row[0])
            avion    = str(row[2]).strip() if row[2] else None
            doc_type = str(row[3]).strip() if row[3] else None
            category = str(row[4]).strip() if row[4] else None
            es_ref   = str(row[5]).strip() if row[5] else None
            ata      = str(row[6]).strip() if row[6] else None

            if doc_type:
                doc_type_norm = TYPE_MAP.get(doc_type.lower(), doc_type.upper())
            else:
                doc_type_norm = None

            result = await db.execute(select(Document).where(Document.id == doc_id))
            doc = result.scalar_one_or_none()
            if not doc:
                skipped += 1
                continue

            changed = False
            if avion and avion != (doc.aircraft_registration or ""):
                doc.aircraft_registration = avion
                changed = True
            if doc_type_norm and doc_type_norm != (doc.doc_type or ""):
                doc.doc_type = doc_type_norm
                changed = True
            if category and category != (doc.category or ""):
                doc.category = category
                changed = True
            if es_ref and es_ref != (doc.es_reference or ""):
                doc.es_reference = es_ref
                changed = True
            if ata and ata != (doc.ata_chapter or ""):
                doc.ata_chapter = ata
                changed = True

            if changed:
                doc.manually_corrected = True
                doc.needs_review = False
                updated += 1

        except Exception as e:
            errors.append(f"Ligne {row[0]}: {str(e)}")

    await db.commit()

    return {
        "message": "Synchronisation terminee",
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:10],
    }


@router.api_route("/{doc_id}/file", methods=["GET", "HEAD"], summary="Telecharger / afficher le PDF original")
async def serve_document_file(
    doc_id: int,
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(Document).where(Document.id == doc_id))
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(404, f"Document #{doc_id} introuvable")

    original_path = (doc.original_path or "").lstrip("\\\\?\\").lstrip("//?/")

    candidate_paths = [
        original_path,
        os.path.join(settings.archive_root_path, original_path),
        os.path.join(settings.upload_path, os.path.basename(original_path)),
        os.path.join(settings.archive_root_path, doc.filename),
    ]

    file_path = None
    for path in candidate_paths:
        if path and os.path.isfile(path):
            file_path = path
            break

    if not file_path:
        raise HTTPException(
            404,
            f"Fichier physique introuvable pour le document #{doc_id}. "
            f"Chemin enregistre: {original_path}"
        )

    return FileResponse(
        path=file_path,
        media_type="application/pdf",
        filename=doc.filename,
        headers={"Content-Disposition": f'inline; filename="{doc.filename}"'},
    )


@router.get("/{doc_id}", response_model=DocumentResponse, summary="Detail document")
async def get_document(doc_id: int, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Document).where(Document.id == doc_id))
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(404, f"Document #{doc_id} introuvable")
    return DocumentResponse.model_validate(doc)


@router.patch("/{doc_id}", response_model=DocumentResponse, summary="Corriger un document")
async def update_document(
    doc_id: int,
    payload: DocumentUpdate,
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(select(Document).where(Document.id == doc_id))
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(404, f"Document #{doc_id} introuvable")

    update_data = payload.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        if hasattr(doc, field):
            setattr(doc, field, value)

    doc.manually_corrected = True
    await db.commit()
    await db.refresh(doc)
    logger.info(f"Document #{doc_id} corrige manuellement: {update_data}")
    return DocumentResponse.model_validate(doc)


@router.delete("/{doc_id}", summary="Supprimer un document")
async def delete_document(doc_id: int, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Document).where(Document.id == doc_id))
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(404, f"Document #{doc_id} introuvable")

    from backend.models.check import Alert
    from sqlalchemy import delete
    await db.execute(delete(Alert).where(Alert.document_id == doc_id))
    await db.delete(doc)
    await db.commit()
    return {"message": f"Document #{doc_id} supprime"}


@router.get("/{doc_id}/ocr-text", summary="Texte OCR brut")
async def get_ocr_text(doc_id: int, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Document).where(Document.id == doc_id))
    doc = result.scalar_one_or_none()
    if not doc:
        raise HTTPException(404)
    return {"document_id": doc_id, "filename": doc.filename, "ocr_text": doc.ocr_text}